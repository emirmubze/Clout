import io
from datetime import datetime
from decimal import Decimal

from django.utils import timezone
from django.utils.timezone import is_aware, localtime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import CustomUser, Order


def _format_datetime(dt):
    """
    Format a datetime object to standard YYYY-MM-DD HH:MM:SS.
    Converts aware datetimes to the current/local timezone.
    """
    if not dt:
        return "N/A"
    try:
        if is_aware(dt):
            dt = localtime(dt)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(dt)


def _get_user_full_name(user):
    """
    Safely get the user's full name, falling back to name or username.
    """
    if not user:
        return "N/A"
    full_name = ""
    if hasattr(user, "display_name") and user.display_name:
        full_name = user.display_name.strip()
    elif hasattr(user, "get_full_name") and user.get_full_name():
        full_name = user.get_full_name().strip()
    elif getattr(user, "name", None):
        full_name = str(user.name).strip()
    elif getattr(user, "username", None):
        full_name = str(user.username).strip()
    return full_name or "N/A"


def _get_course_access_status(user):
    """
    Determine course access status independently from payment status.
    """
    if user and getattr(user, "course_access_approved", False):
        return "APPROVED"
    return "NOT APPROVED"


def _apply_excel_styling(ws, headers, rows_count):
    """
    Apply professional styling to an openpyxl worksheet:
    - Styled bold header row with emerald green background and white text.
    - Freeze panes at row 2 so headers remain visible when scrolling.
    - Add auto-filters to the header row.
    - Zebra striping or clean border formatting for data rows.
    - Auto-fit column widths with protective minimum padding.
    """
    header_fill = PatternFill(
        start_color="059669",
        end_color="059669",
        fill_type="solid",
    )
    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color="FFFFFF",
    )
    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=False,
    )

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    ws.freeze_panes = "A2"

    last_col_letter = get_column_letter(len(headers))
    last_row = max(rows_count + 1, 2)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center")
    number_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, rows_count + 2):
        ws.row_dimensions[row].height = 20
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = data_font
            cell.border = thin_border

            header_name = headers[col_idx - 1]
            if header_name in ("User ID", "Payment Status", "Course Access Status", "Currency"):
                cell.alignment = center_alignment
            elif header_name == "Payment Amount" and isinstance(cell.value, (int, float, Decimal)):
                cell.alignment = number_alignment
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = data_alignment

    for col in ws.columns:
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                cell_str = f"{val:.2f}" if isinstance(val, (int, float, Decimal)) else str(val)
                max_len = max(max_len, len(cell_str))
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)


def generate_users_workbook(export_type):
    """
    Generate an in-memory openpyxl Workbook for the requested export type:
    - 'paid': Users who have a verified successful payment (Order.paid=True).
    - 'unpaid': Registered users without a verified successful payment.
    - 'all': All registered users with clear PAID/UNPAID status.

    Returns:
        tuple: (workbook, count, filename)
    """
    export_type = (export_type or "").strip().lower()
    today_str = timezone.now().strftime("%Y-%m-%d")

    if export_type == "paid":
        return _build_paid_users_workbook(today_str)
    elif export_type == "unpaid":
        return _build_unpaid_users_workbook(today_str)
    elif export_type == "all":
        return _build_all_users_workbook(today_str)
    else:
        raise ValueError(f"Unknown export type: '{export_type}'. Expected 'paid', 'unpaid', or 'all'.")


def _build_paid_users_workbook(today_str):
    """
    Export only users who have a verified successful payment.
    """
    headers = [
        "User ID",
        "Full Name",
        "Email",
        "Phone Number",
        "Registration Date",
        "Payment Status",
        "Payment Amount",
        "Currency",
        "Razorpay Order ID",
        "Razorpay Payment ID",
        "Payment Date",
        "Course Access Status",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paid Users"
    ws.append(headers)

    # Fetch verified orders linked to users, newest first
    paid_orders = (
        Order.objects.filter(paid=True, user__isnull=False)
        .select_related("user")
        .order_by("-created_at", "-id")
    )

    seen_user_ids = set()
    rows_written = 0

    for order in paid_orders:
        user = order.user
        if not user or user.id in seen_user_ids:
            continue
        seen_user_ids.add(user.id)

        try:
            amount_val = float(order.amount)
        except (ValueError, TypeError):
            amount_val = order.amount or 0.0

        row_data = [
            user.id,
            _get_user_full_name(user),
            user.email or "N/A",
            user.phone_number or "N/A",
            _format_datetime(user.date_joined),
            "PAID",
            amount_val,
            order.currency or "N/A",
            order.razorpay_order_id or "N/A",
            order.razorpay_payment_id or "N/A",
            _format_datetime(order.created_at),
            _get_course_access_status(user),
        ]
        ws.append(row_data)
        rows_written += 1

    _apply_excel_styling(ws, headers, rows_written)
    filename = f"paid_users_{today_str}.xlsx"
    return wb, rows_written, filename


def _build_unpaid_users_workbook(today_str):
    """
    Export registered users who have NOT completed a verified successful payment.
    """
    headers = [
        "User ID",
        "Full Name",
        "Email",
        "Phone Number",
        "Registration Date",
        "Payment Status",
        "Course Access Status",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Unpaid Users"
    ws.append(headers)

    paid_user_ids = set(
        Order.objects.filter(paid=True, user__isnull=False)
        .values_list("user_id", flat=True)
    )

    unpaid_users = (
        CustomUser.objects.exclude(id__in=paid_user_ids)
        .order_by("id")
    )

    rows_written = 0
    for user in unpaid_users:
        row_data = [
            user.id,
            _get_user_full_name(user),
            user.email or "N/A",
            user.phone_number or "N/A",
            _format_datetime(user.date_joined),
            "UNPAID",
            _get_course_access_status(user),
        ]
        ws.append(row_data)
        rows_written += 1

    _apply_excel_styling(ws, headers, rows_written)
    filename = f"unpaid_users_{today_str}.xlsx"
    return wb, rows_written, filename


def _build_all_users_workbook(today_str):
    """
    Export all registered users with verified payment and course access status.
    """
    headers = [
        "User ID",
        "Full Name",
        "Email",
        "Phone Number",
        "Registration Date",
        "Payment Status",
        "Payment Amount",
        "Currency",
        "Razorpay Order ID",
        "Razorpay Payment ID",
        "Payment Date",
        "Course Access Status",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Users"
    ws.append(headers)

    # Build a map of user_id -> latest verified Order
    paid_orders = (
        Order.objects.filter(paid=True, user__isnull=False)
        .order_by("-created_at", "-id")
    )
    user_order_map = {}
    for order in paid_orders:
        if order.user_id not in user_order_map:
            user_order_map[order.user_id] = order

    users = CustomUser.objects.all().order_by("id")
    rows_written = 0

    for user in users:
        order = user_order_map.get(user.id)
        if order:
            payment_status = "PAID"
            try:
                amount_val = float(order.amount)
            except (ValueError, TypeError):
                amount_val = order.amount or 0.0
            currency_val = order.currency or "N/A"
            order_id_val = order.razorpay_order_id or "N/A"
            payment_id_val = order.razorpay_payment_id or "N/A"
            payment_date_val = _format_datetime(order.created_at)
        else:
            payment_status = "UNPAID"
            amount_val = "N/A"
            currency_val = "N/A"
            order_id_val = "N/A"
            payment_id_val = "N/A"
            payment_date_val = "N/A"

        row_data = [
            user.id,
            _get_user_full_name(user),
            user.email or "N/A",
            user.phone_number or "N/A",
            _format_datetime(user.date_joined),
            payment_status,
            amount_val,
            currency_val,
            order_id_val,
            payment_id_val,
            payment_date_val,
            _get_course_access_status(user),
        ]
        ws.append(row_data)
        rows_written += 1

    _apply_excel_styling(ws, headers, rows_written)
    filename = f"all_users_{today_str}.xlsx"
    return wb, rows_written, filename


def export_users_to_bytes(export_type):
    """
    Helper that generates the workbook and saves it into an in-memory BytesIO buffer.
    Returns:
        tuple: (bytes_data, count, filename)
    """
    wb, count, filename = generate_users_workbook(export_type)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue(), count, filename
